from pathlib import Path
from docx import Document

BASE_DIR = Path(__file__).resolve().parent.parent
DOCX = BASE_DIR / "web" / "assets" / "report" / "bao-cao-brain-tumor-v2.docx"


def _rows(table):
    return [[c.text.strip() for c in r.cells] for r in table.rows]


def test_table1_has_b4b5_and_measured_sizes():
    assert DOCX.exists(), f"File docx khong ton tai: {DOCX}"
    doc = Document(DOCX)
    assert len(doc.tables) >= 2, "Bao cao can co it nhat 2 bang"
    rows = _rows(doc.tables[0])

    assert len(rows) == 5, f"Bang 1 can co 5 hang (1 header + 4 data), nhan duoc {len(rows)}"

    best = next((r for r in rows[1:] if "Block 4" in r[0]), None)
    assert best is not None, "Thieu dong VGG16 Block 4+5 trong Bang 1"
    assert best[1:] == [
        "14.716.740",
        "155.3 MB",
        "95.12%",
        "0.9535",
        "0.9512",
        "0.9502",
        "75.42 ms",
    ], f"Chi so cua VGG16 Block 4+5 khong dung: {best}"

    block5 = next((r for r in rows[1:] if "Block 5" in r[0] and "Block 4" not in r[0]), None)
    assert block5 is not None, "Thieu dong VGG16 Block 5 trong Bang 1"
    assert block5[2] == "110.3 MB", f"Dung luong VGG16 Block 5 phai la 110.3 MB, nhan duoc: {block5[2]}"

    all_cells_text = [cell for r in rows for cell in r]
    assert "58.0 MB" not in all_cells_text, "Bang 1 van con gia tri cu 58.0 MB"


def test_ablation_uses_measured_trainable_params():
    doc = Document(DOCX)
    cells = [c.text.strip() for r in doc.tables[1].rows for c in r.cells]
    text = " ".join(cells)
    assert "7.081.476" in text, "Bang 2 (Ablation) phai chua 7.081.476 tham so tinh chinh cua Block 5"
    assert "7.079.428" not in text, "Bang 2 van con gia tri cu 7.079.428"
